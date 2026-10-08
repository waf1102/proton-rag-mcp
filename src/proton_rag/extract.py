"""MIME/attachment parsing in a disposable process; raw input never staged on disk."""

import io
import os
import multiprocessing
import resource
import zipfile
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from .config import Settings

MAX_RAW = Settings().max_message_bytes
MAX_TEXT = Settings().max_text_chars
MAX_PARTS = Settings().max_parts


class HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.hidden += 1

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def parse(raw, settings=None):
    settings = settings or Settings()
    if len(raw) > settings.max_message_bytes:
        return {"text": "", "skipped": ["message_size"]}
    msg = BytesParser(policy=policy.default).parsebytes(raw)
    texts, skipped = [], []
    parts = list(msg.walk())
    if len(parts) > settings.max_parts:
        return {"text": "", "skipped": ["part_count"]}
    for part in parts:
        if part.is_multipart():
            continue
        kind = part.get_content_type()
        payload = part.get_payload(decode=True) or b""
        try:
            if kind in ("text/plain", "text/csv"):
                text = payload.decode(part.get_content_charset() or "utf-8", errors="replace")
            elif kind == "text/html" and part.get_content_disposition() != "attachment":
                html = HTMLText()
                html.feed(payload.decode(part.get_content_charset() or "utf-8", errors="replace"))
                text = " ".join(html.parts)
            elif kind == "application/pdf":
                from pypdf import PdfReader

                reader = PdfReader(io.BytesIO(payload))
                if len(reader.pages) > 100:
                    raise ValueError("page bound")
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
            elif kind == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                from docx import Document

                with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                    if sum(x.file_size for x in archive.infolist()) > 16 * 1024 * 1024:
                        raise ValueError("expanded size bound")
                doc = Document(io.BytesIO(payload))
                text = "\n".join(p.text for p in doc.paragraphs)
                text += "\n" + "\n".join(
                    " | ".join(c.text for c in r.cells) for table in doc.tables for r in table.rows
                )
            else:
                skipped.append("unsupported")
                continue
            texts.append(text[: settings.max_text_chars])
        except Exception:
            skipped.append("malformed_attachment")
    if msg.defects:
        skipped.append("mime_defects")
    metadata = {
        name: str(msg.get(header, ""))
        for name, header in [
            ("subject", "Subject"),
            ("sender", "From"),
            ("recipients", "To"),
            ("cc", "Cc"),
            ("date", "Date"),
            ("message_id", "Message-ID"),
        ]
    }
    metadata["attachments"] = [str(p.get_filename()) for p in parts if p.get_filename()]
    body = "\n\n".join(texts)[: settings.max_text_chars].strip()
    headers = "\n".join(f"{k}: {v}" for k, v in metadata.items() if v)
    return {
        "text": (headers + "\n\n" + body).strip() if body else "",
        "metadata": metadata,
        "skipped": skipped,
    }


def _worker(raw, pipe, settings):
    try:
        # Third-party parsers must not put private document fragments in diagnostics.
        sink = os.open(os.devnull, os.O_WRONLY)
        os.dup2(sink, 1)
        os.dup2(sink, 2)
        os.close(sink)
        resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024,) * 2)
        resource.setrlimit(resource.RLIMIT_CPU, (5, 5))
        pipe.send(parse(raw, settings))
    except BaseException:
        pipe.send({"text": "", "skipped": ["parse_failed"]})
    finally:
        pipe.close()


def extract(raw, timeout=None, settings=None):
    settings = settings or Settings()
    timeout = timeout or settings.parser_timeout
    ctx = multiprocessing.get_context("spawn")
    receive, send = ctx.Pipe(duplex=False)
    process = ctx.Process(target=_worker, args=(raw, send, settings), daemon=True)
    process.start()
    send.close()
    try:
        if receive.poll(timeout):
            try:
                return receive.recv()
            except EOFError:
                pass
        return {"text": "", "skipped": ["parse_timeout_or_resource_limit"]}
    finally:
        if process.is_alive():
            process.terminate()
        process.join()
        receive.close()
